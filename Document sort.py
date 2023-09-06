import os
import shutil
import PyPDF4
import docx
from pptx import Presentation
import openpyxl
from PyPDF4.utils import PdfReadError

# Function to sort files by format
def sort_files(folder_path):
    # Create folders to store different file formats
    format_folders = {
        'pdf': 'PDF',
        'doc': 'Word',
        'docx': 'Word',
        'ppt': 'PowerPoint',
        'pptx': 'PowerPoint',
        'xls': 'Excel',
        'xlsx': 'Excel'
    }

    # Create folders for different formats
    for format_folder in set(format_folders.values()):
        os.makedirs(os.path.join(folder_path, format_folder), exist_ok=True)

    # Create folder for corrupted files
    corrupted_folder = os.path.join(folder_path, 'Corrupted Files')
    os.makedirs(corrupted_folder, exist_ok=True)

    # Iterate through files in the folder
    for file_name in os.listdir(folder_path):
        file_path = os.path.join(folder_path, file_name)

        # Skip directories and already sorted folders
        if os.path.isdir(file_path) or file_name in format_folders.values() or file_name == 'Corrupted Files':
            continue

        # Determine file format
        file_extension = os.path.splitext(file_name)[1][1:].lower()

        # Check for corruption
        is_corrupted = is_file_corrupted(file_path, file_extension)

        # Move file to the respective format folder or corrupted folder
        if file_extension in format_folders and not is_corrupted:
            format_folder = format_folders[file_extension]
            new_file_path = os.path.join(folder_path, format_folder, file_name)
            shutil.move(file_path, new_file_path)
        else:
            new_file_path = os.path.join(corrupted_folder, file_name)
            shutil.move(file_path, new_file_path)

# Function to check if a PDF file is corrupted
def is_pdf_corrupted(file_path):
    try:
        with open(file_path, 'rb') as file:
            reader = PyPDF4.PdfFileReader(file)
            reader.getNumPages()  # Access a property to check if the file is valid
        return False
    except (PdfReadError, FileNotFoundError):
        return True

# Function to check if a Word file is corrupted
def is_word_corrupted(file_path):
    try:
        doc = docx.Document(file_path)
        return False
    except (docx.exceptions.PackageNotFoundError, docx.exceptions.InvalidFileException):
        return True

# Function to check if a PowerPoint file is corrupted
def is_powerpoint_corrupted(file_path):
    try:
        presentation = Presentation(file_path)
        return False
    except pptx.exc.PackageNotFoundError:
        return True

# Function to check if an Excel file is corrupted
def is_excel_corrupted(file_path):
    try:
        workbook = openpyxl.load_workbook(file_path)
        return False
    except openpyxl.utils.exceptions.InvalidFileException:
        return True

# Function to check if a file is corrupted
def is_file_corrupted(file_path, file_extension):
    if file_extension == 'pdf':
        return is_pdf_corrupted(file_path)
    elif file_extension in ('doc', 'docx'):
        return is_word_corrupted(file_path)
    elif file_extension in ('ppt', 'pptx'):
        return is_powerpoint_corrupted(file_path)
    elif file_extension in ('xls', 'xlsx'):
        return is_excel_corrupted(file_path)
    else:
        return True  # Treat unknown file formats as corrupted

# Example usage
folder_path = r'D:\Documents'
sort_files(folder_path)
